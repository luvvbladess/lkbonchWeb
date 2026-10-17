import csv
import io
import logging
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook

from app.errors import ApiError, invalid
from app.schemas import ImportBatch, ImportIssue
from app.services.common import many, must, one, parse_amount, university_zone
from app.services.desk import read_stored, store_bytes

logger = logging.getLogger("lkbonch")

DOMAINS = (
    "students",
    "employees",
    "groups",
    "disciplines",
    "elective_blocks",
    "curriculum",
    "lessons",
    "scholarships",
    "contracts",
    "ledger",
)

REQUIRED_COLUMNS = {
    "students": ["last_name", "first_name", "group_name", "entry_year", "course"],
    "employees": ["last_name", "first_name", "login"],
    "groups": ["name", "course", "entry_year", "org_unit"],
    "disciplines": ["code", "name"],
    "elective_blocks": ["group_name", "entry_year", "semester_title", "title"],
    "curriculum": ["group_name", "entry_year", "discipline_code", "semester_title", "kind"],
    "lessons": ["discipline_code", "group_name", "entry_year"],
    "scholarships": ["student_card_number", "title", "amount"],
    "contracts": ["student_card_number", "kind", "number"],
    "ledger": ["contract_number", "direction", "amount", "occurred_on"],
}


class RowProblem(Exception):
    def __init__(self, field: str, message: str) -> None:
        super().__init__(message)
        self.field = field
        self.message = message


class AlreadyThere(Exception):
    pass


def expand_week_dates(starts_on: date, ends_on: date, weekday: int, week: str) -> list[date]:
    if weekday < 1 or weekday > 7:
        raise ValueError("weekday")
    if week not in ("all", "odd", "even"):
        raise ValueError("week")
    found = []
    day = starts_on
    while day <= ends_on:
        if day.isoweekday() == weekday:
            number = day.isocalendar().week
            if week == "all" or (week == "odd" and number % 2 == 1) or (week == "even" and number % 2 == 0):
                found.append(day)
        day = date.fromordinal(day.toordinal() + 1)
    return found


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows_from_xlsx(data: bytes) -> list[tuple[int, dict[str, str]]]:
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = workbook.active
    matrix = list(sheet.iter_rows(values_only=True))
    workbook.close()
    return _matrix(matrix)


def _rows_from_csv(data: bytes) -> list[tuple[int, dict[str, str]]]:
    text = data.decode("utf-8-sig")
    try:
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    matrix = list(csv.reader(io.StringIO(text), dialect))
    return _matrix(matrix)


def _matrix(matrix) -> list[tuple[int, dict[str, str]]]:
    if not matrix:
        return []
    headers = [_text(cell).lower().replace(" ", "_") for cell in matrix[0]]
    rows = []
    for offset, raw in enumerate(matrix[1:], start=2):
        mapped = {headers[index]: _text(raw[index] if index < len(raw) else "") for index in range(len(headers)) if headers[index]}
        if any(mapped.values()):
            rows.append((offset, mapped))
    return rows


def _parse(filename: str, data: bytes):
    lower = filename.lower()
    if lower.endswith(".xlsx"):
        return _rows_from_xlsx(data)
    return _rows_from_csv(data)


def _need(raw: dict, field: str) -> str:
    value = raw.get(field, "").strip()
    if not value:
        raise RowProblem(field, "Поле не заполнено")
    return value


def _int(raw: dict, field: str) -> int:
    try:
        return int(_need(raw, field))
    except ValueError as exc:
        raise RowProblem(field, "Ожидается целое число") from exc


def _date(value: str, field: str) -> date | None:
    if not value:
        return None
    text = value.strip()
    for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    raise RowProblem(field, "Дата в формате ГГГГ-ММ-ДД")


def _time(value: str, field: str) -> time:
    text = value.strip()
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    raise RowProblem(field, "Время в формате ЧЧ:ММ")


def _one_group(conn, name: str, year: int):
    rows = many(
        conn,
        "SELECT id FROM groups WHERE lower(name) = lower(%s) AND entry_year = %s",
        (name, year),
    )
    if not rows:
        raise RowProblem("group_name", "Группа не найдена")
    if len(rows) > 1:
        raise RowProblem("group_name", "Несколько групп с таким названием и годом поступления")
    return rows[0]["id"]


def _one_org(conn, name: str):
    rows = many(conn, "SELECT id FROM org_units WHERE lower(name) = lower(%s)", (name,))
    if not rows:
        raise RowProblem("org_unit", "Подразделение не найдено")
    if len(rows) > 1:
        raise RowProblem("org_unit", "Несколько подразделений с таким названием")
    return rows[0]["id"]


def _one_semester(conn, title: str):
    rows = many(conn, "SELECT * FROM semesters WHERE lower(title) = lower(%s)", (title,))
    if not rows:
        raise RowProblem("semester_title", "Семестр не найден")
    if len(rows) > 1:
        raise RowProblem("semester_title", "Несколько семестров с таким названием")
    return rows[0]


def _active_semester(conn):
    row = one(conn, "SELECT * FROM semesters WHERE status = 'active'")
    if row is None:
        raise RowProblem("semester_title", "Нет активного семестра, укажите semester_title")
    return row


def _student_by_card(conn, card: str):
    if not card:
        return None
    return one(
        conn,
        """
        SELECT u.id FROM users u
        JOIN student_profiles sp ON sp.user_id = u.id
        WHERE sp.student_card_number = %s
        """,
        (card,),
    )


def _discipline_id(conn, code: str):
    row = one(conn, "SELECT id FROM disciplines WHERE lower(code) = lower(%s)", (code,))
    if row is None:
        raise RowProblem("discipline_code", "Дисциплина не найдена")
    return row["id"]


def _teacher(conn, login: str):
    if not login:
        return None
    row = one(conn, "SELECT id, role FROM users WHERE login = %s", (login,))
    if row is None or row["role"] != "employee":
        raise RowProblem("teacher_login", "Преподаватель с таким логином не найден")
    return row["id"]


def _import_students(conn, raw, mode, write):
    last_name = _need(raw, "last_name")
    first_name = _need(raw, "first_name")
    middle = raw.get("middle_name", "")
    group_id = _one_group(conn, _need(raw, "group_name"), _int(raw, "entry_year"))
    course = _int(raw, "course")
    if course < 1 or course > 6:
        raise RowProblem("course", "Курс от 1 до 6")
    funding = raw.get("funding") or "budget"
    if funding not in ("budget", "paid", "targeted"):
        raise RowProblem("funding", "Форма обучения: budget, paid или targeted")
    card = raw.get("student_card_number", "")
    book = raw.get("record_book_number", "")
    email = raw.get("email", "") or None
    phone = raw.get("phone", "") or None
    if not email and not phone:
        raise RowProblem("email", "Нужна почта или телефон")
    if not card and not book:
        raise RowProblem("student_card_number", "Нужен номер студенческого или зачётки")
    by_card = _student_by_card(conn, card)
    by_book = None
    if book:
        by_book = one(
            conn,
            """
            SELECT u.id FROM users u
            JOIN student_profiles sp ON sp.user_id = u.id
            WHERE sp.record_book_number = %s
            """,
            (book,),
        )
    if by_card and by_book and by_card["id"] != by_book["id"]:
        raise RowProblem("record_book_number", "Студенческий и зачётка указывают на разных людей")
    existing = by_card or by_book
    if existing and mode == "insert_only":
        raise RowProblem("student_card_number", "Такой студент уже есть")
    if not write:
        return
    if existing:
        conn.execute(
            """
            UPDATE users
            SET last_name = %s, first_name = %s, middle_name = %s, email = COALESCE(%s, email), phone = COALESCE(%s, phone)
            WHERE id = %s
            """,
            (last_name, first_name, middle, email, phone, existing["id"]),
        )
        conn.execute(
            """
            UPDATE student_profiles
            SET group_id = %s, course = %s, funding = %s,
                record_book_number = NULLIF(%s, ''), student_card_number = NULLIF(%s, '')
            WHERE user_id = %s
            """,
            (group_id, course, funding, book, card, existing["id"]),
        )
        return
    created = conn.execute(
        """
        INSERT INTO users (role, status, last_name, first_name, middle_name, email, phone)
        VALUES ('student', 'pending', %s, %s, %s, %s, %s)
        RETURNING id
        """,
        (last_name, first_name, middle, email, phone),
    ).fetchone()
    conn.execute(
        """
        INSERT INTO student_profiles (
            user_id, group_id, course, funding, record_book_number, student_card_number
        )
        VALUES (%s, %s, %s, %s, NULLIF(%s, ''), NULLIF(%s, ''))
        """,
        (created["id"], group_id, course, funding, book, card),
    )


def _import_employees(conn, raw, mode, write):
    login = _need(raw, "login")
    existing = one(conn, "SELECT id, role FROM users WHERE login = %s", (login,))
    if existing and existing["role"] != "employee":
        raise RowProblem("login", "Этот логин принадлежит другой роли")
    if existing and mode == "insert_only":
        raise RowProblem("login", "Сотрудник с таким логином уже есть")
    org_id = _one_org(conn, raw["org_unit"]) if raw.get("org_unit") else None
    if not write:
        return
    if existing:
        conn.execute(
            """
            UPDATE users
            SET last_name = %s, first_name = %s, middle_name = %s,
                email = NULLIF(%s, ''), phone = NULLIF(%s, '')
            WHERE id = %s
            """,
            (_need(raw, "last_name"), _need(raw, "first_name"), raw.get("middle_name", ""), raw.get("email", ""), raw.get("phone", ""), existing["id"]),
        )
        conn.execute(
            """
            INSERT INTO employee_profiles (user_id, org_unit_id, position)
            VALUES (%s, %s, %s)
            ON CONFLICT (user_id) DO UPDATE
            SET org_unit_id = EXCLUDED.org_unit_id, position = EXCLUDED.position
            """,
            (existing["id"], org_id, raw.get("position", "")),
        )
        return
    created = conn.execute(
        """
        INSERT INTO users (role, status, last_name, first_name, middle_name, login, email, phone)
        VALUES ('employee', 'pending', %s, %s, %s, %s, NULLIF(%s, ''), NULLIF(%s, ''))
        RETURNING id
        """,
        (_need(raw, "last_name"), _need(raw, "first_name"), raw.get("middle_name", ""), login, raw.get("email", ""), raw.get("phone", "")),
    ).fetchone()
    conn.execute(
        "INSERT INTO employee_profiles (user_id, org_unit_id, position) VALUES (%s, %s, %s)",
        (created["id"], org_id, raw.get("position", "")),
    )


def _import_groups(conn, raw, mode, write):
    name = _need(raw, "name")
    year = _int(raw, "entry_year")
    course = _int(raw, "course")
    org_id = _one_org(conn, _need(raw, "org_unit"))
    existing = one(
        conn,
        "SELECT id FROM groups WHERE org_unit_id = %s AND lower(name) = lower(%s) AND entry_year = %s",
        (org_id, name, year),
    )
    if existing and mode == "insert_only":
        raise RowProblem("name", "Такая группа уже есть")
    if not write:
        return
    if existing:
        conn.execute("UPDATE groups SET course = %s WHERE id = %s", (course, existing["id"]))
        return
    conn.execute(
        "INSERT INTO groups (org_unit_id, name, course, entry_year) VALUES (%s, %s, %s, %s)",
        (org_id, name, course, year),
    )


def _import_disciplines(conn, raw, mode, write):
    code = _need(raw, "code")
    existing = one(conn, "SELECT id FROM disciplines WHERE lower(code) = lower(%s)", (code,))
    if existing and mode == "insert_only":
        raise RowProblem("code", "Дисциплина с таким кодом уже есть")
    org_id = _one_org(conn, raw["org_unit"]) if raw.get("org_unit") else None
    if not write:
        return
    if existing:
        conn.execute(
            "UPDATE disciplines SET name = %s, org_unit_id = %s WHERE id = %s",
            (_need(raw, "name"), org_id, existing["id"]),
        )
        return
    conn.execute(
        "INSERT INTO disciplines (code, name, org_unit_id) VALUES (%s, %s, %s)",
        (code, _need(raw, "name"), org_id),
    )


def _import_blocks(conn, raw, mode, write):
    group_id = _one_group(conn, _need(raw, "group_name"), _int(raw, "entry_year"))
    semester = _one_semester(conn, _need(raw, "semester_title"))
    title = _need(raw, "title")
    existing = one(
        conn,
        """
        SELECT id FROM elective_blocks
        WHERE group_id = %s AND semester_id = %s AND lower(title) = lower(%s)
        """,
        (group_id, semester["id"], title),
    )
    if existing and mode == "insert_only":
        raise RowProblem("title", "Такой блок уже есть")
    if not write or existing:
        return
    conn.execute(
        "INSERT INTO elective_blocks (group_id, semester_id, title) VALUES (%s, %s, %s)",
        (group_id, semester["id"], title),
    )


def _import_curriculum(conn, raw, mode, write):
    group_id = _one_group(conn, _need(raw, "group_name"), _int(raw, "entry_year"))
    discipline_id = _discipline_id(conn, _need(raw, "discipline_code"))
    semester = _one_semester(conn, _need(raw, "semester_title"))
    kind = _need(raw, "kind")
    if kind not in ("required", "elective"):
        raise RowProblem("kind", "Вид: required или elective")
    block_id = None
    if kind == "elective":
        title = _need(raw, "block_title")
        block = one(
            conn,
            """
            SELECT id FROM elective_blocks
            WHERE group_id = %s AND semester_id = %s AND lower(title) = lower(%s)
            """,
            (group_id, semester["id"], title),
        )
        if block is None:
            raise RowProblem("block_title", "Блок выборных не найден")
        block_id = block["id"]
    teacher_id = _teacher(conn, raw.get("teacher_login", ""))
    try:
        credits = Decimal(raw["credits"]) if raw.get("credits") else Decimal("0")
    except InvalidOperation as exc:
        raise RowProblem("credits", "Кредиты записаны не числом") from exc
    existing = one(
        conn,
        """
        SELECT id FROM group_disciplines
        WHERE group_id = %s AND discipline_id = %s AND semester_id = %s
        """,
        (group_id, discipline_id, semester["id"]),
    )
    if existing and mode == "insert_only":
        raise RowProblem("discipline_code", "Дисциплина уже стоит в плане этой группы")
    if not write:
        return
    if existing:
        conn.execute(
            """
            UPDATE group_disciplines
            SET kind = %s, elective_block_id = %s, credits = %s, teacher_user_id = %s
            WHERE id = %s
            """,
            (kind, block_id, credits, teacher_id, existing["id"]),
        )
        return
    conn.execute(
        """
        INSERT INTO group_disciplines (
            group_id, discipline_id, semester_id, kind, elective_block_id, credits, teacher_user_id
        )
        VALUES (%s, %s, %s, %s, %s, %s, %s)
        """,
        (group_id, discipline_id, semester["id"], kind, block_id, credits, teacher_id),
    )


def _lesson_moments(conn, raw, semester):
    zone = university_zone(conn)
    if raw.get("date"):
        day = _date(raw["date"], "date")
        start = _time(raw.get("starts_at") or raw.get("start_time", ""), "starts_at")
        end = _time(raw.get("ends_at") or raw.get("end_time", ""), "ends_at")
        return [(datetime.combine(day, start, zone), datetime.combine(day, end, zone))]
    if raw.get("weekday"):
        week = (raw.get("week") or "all").lower()
        try:
            days = expand_week_dates(semester["starts_on"], semester["ends_on"], _int(raw, "weekday"), week)
        except ValueError as exc:
            raise RowProblem("weekday", "День недели от 1 до 7, неделя all, odd или even") from exc
        if not days:
            raise RowProblem("weekday", "В семестре нет таких дат")
        start = _time(_need(raw, "start_time"), "start_time")
        end = _time(_need(raw, "end_time"), "end_time")
        return [(datetime.combine(day, start, zone), datetime.combine(day, end, zone)) for day in days]
    raise RowProblem("date", "Укажите дату пары или день недели")


def _import_lessons(conn, raw, mode, write):
    discipline_id = _discipline_id(conn, _need(raw, "discipline_code"))
    group_id = _one_group(conn, _need(raw, "group_name"), _int(raw, "entry_year"))
    semester = _one_semester(conn, raw["semester_title"]) if raw.get("semester_title") else _active_semester(conn)
    teacher_id = _teacher(conn, raw.get("teacher_login", ""))
    room = raw.get("room", "")
    spans = _lesson_moments(conn, raw, semester)
    for starts, ends in spans:
        if ends <= starts:
            raise RowProblem("ends_at", "Конец пары раньше начала")
        existing = one(
            conn,
            """
            SELECT l.id
            FROM lessons l
            JOIN lesson_groups lg ON lg.lesson_id = l.id
            WHERE lg.group_id = %s AND l.discipline_id = %s AND l.starts_at = %s
            """,
            (group_id, discipline_id, starts),
        )
        if existing and mode == "insert_only":
            raise RowProblem("starts_at", "Такая пара уже есть")
        if not write:
            continue
        if existing:
            conn.execute(
                """
                UPDATE lessons
                SET ends_at = %s, room = %s, teacher_user_id = %s, semester_id = %s
                WHERE id = %s
                """,
                (ends, room, teacher_id, semester["id"], existing["id"]),
            )
            continue
        created = conn.execute(
            """
            INSERT INTO lessons (
                semester_id, discipline_id, teacher_user_id, starts_at, ends_at, room,
                publication_status, published_at
            )
            VALUES (%s, %s, %s, %s, %s, %s, 'published', now())
            RETURNING id
            """,
            (semester["id"], discipline_id, teacher_id, starts, ends, room),
        ).fetchone()
        conn.execute(
            "INSERT INTO lesson_groups (lesson_id, group_id) VALUES (%s, %s)",
            (created["id"], group_id),
        )


def _import_scholarships(conn, raw, mode, write):
    student = _student_by_card(conn, _need(raw, "student_card_number"))
    if student is None:
        raise RowProblem("student_card_number", "Студент не найден")
    title = _need(raw, "title")
    amount = parse_amount(raw.get("amount"), allow_zero=True)
    payout = _date(raw.get("next_payout_on", ""), "next_payout_on")
    existing = one(
        conn,
        """
        SELECT id FROM scholarship_grants
        WHERE student_user_id = %s AND lower(title) = lower(%s)
        ORDER BY created_at DESC LIMIT 1
        """,
        (student["id"], title),
    )
    if existing and mode == "insert_only":
        raise RowProblem("title", "Такая стипендия уже есть")
    if not write:
        return
    if existing:
        conn.execute(
            "UPDATE scholarship_grants SET amount = %s, next_payout_on = %s, active = true WHERE id = %s",
            (amount, payout, existing["id"]),
        )
        return
    conn.execute(
        """
        INSERT INTO scholarship_grants (student_user_id, title, amount, next_payout_on)
        VALUES (%s, %s, %s, %s)
        """,
        (student["id"], title, amount, payout),
    )


def _import_contracts(conn, raw, mode, write):
    student = _student_by_card(conn, _need(raw, "student_card_number"))
    if student is None:
        raise RowProblem("student_card_number", "Студент не найден")
    kind = _need(raw, "kind")
    if kind not in ("tuition", "housing", "other"):
        raise RowProblem("kind", "Вид договора: tuition, housing или other")
    number = _need(raw, "number")
    existing = one(
        conn,
        "SELECT id FROM contracts WHERE student_user_id = %s AND number = %s",
        (student["id"], number),
    )
    if existing and mode == "insert_only":
        raise RowProblem("number", "Договор с таким номером уже есть")
    if not write:
        return
    opened = _date(raw.get("opened_on", ""), "opened_on")
    if existing:
        conn.execute(
            "UPDATE contracts SET kind = %s, title = %s, opened_on = %s WHERE id = %s",
            (kind, raw.get("title", ""), opened, existing["id"]),
        )
        return
    conn.execute(
        """
        INSERT INTO contracts (student_user_id, kind, number, title, opened_on)
        VALUES (%s, %s, %s, %s, %s)
        """,
        (student["id"], kind, number, raw.get("title", ""), opened),
    )


def _import_ledger(conn, raw, mode, write):
    number = _need(raw, "contract_number")
    found = many(conn, "SELECT id FROM contracts WHERE number = %s", (number,))
    if not found:
        raise RowProblem("contract_number", "Договор не найден")
    if len(found) > 1:
        raise RowProblem("contract_number", "Номер договора встречается у нескольких студентов")
    direction = _need(raw, "direction")
    if direction not in ("charge", "payment"):
        raise RowProblem("direction", "Направление: charge или payment")
    amount = parse_amount(raw.get("amount"))
    occurred = _date(_need(raw, "occurred_on"), "occurred_on")
    description = raw.get("description", "")
    duplicate = one(
        conn,
        """
        SELECT 1 FROM ledger_entries
        WHERE contract_id = %s AND direction = %s AND amount = %s
          AND occurred_on = %s AND description = %s AND reverses_entry_id IS NULL
        """,
        (found[0]["id"], direction, amount, occurred, description),
    )
    if duplicate:
        raise AlreadyThere()
    if not write:
        return
    conn.execute(
        """
        INSERT INTO ledger_entries (contract_id, direction, amount, occurred_on, description)
        VALUES (%s, %s, %s, %s, %s)
        """,
        (found[0]["id"], direction, amount, occurred, description),
    )


HANDLERS = {
    "students": _import_students,
    "employees": _import_employees,
    "groups": _import_groups,
    "disciplines": _import_disciplines,
    "elective_blocks": _import_blocks,
    "curriculum": _import_curriculum,
    "lessons": _import_lessons,
    "scholarships": _import_scholarships,
    "contracts": _import_contracts,
    "ledger": _import_ledger,
}


def _check_headers(domain: str, rows: list[tuple[int, dict]]) -> None:
    if not rows:
        raise invalid("import_empty", "В файле нет строк")
    headers = set(rows[0][1])
    missing = [column for column in REQUIRED_COLUMNS[domain] if column not in headers]
    if domain == "students" and "email" not in headers and "phone" not in headers:
        missing.append("email или phone")
    if missing:
        raise invalid("import_columns", "В файле нет колонок: " + ", ".join(missing))


def _process(conn, domain: str, mode: str, rows, write: bool):
    handler = HANDLERS[domain]
    accepted = 0
    rejected = 0
    issues: list[ImportIssue] = []
    for number, raw in rows:
        try:
            with conn.transaction():
                handler(conn, raw, mode, write)
            accepted += 1
        except AlreadyThere:
            accepted += 1
        except RowProblem as exc:
            rejected += 1
            if len(issues) < 500:
                issues.append(ImportIssue(row_number=number, field_name=exc.field, message=exc.message))
        except ApiError as exc:
            rejected += 1
            if len(issues) < 500:
                issues.append(ImportIssue(row_number=number, field_name="", message=exc.message))
        except Exception:
            logger.exception("import row %s failed", number)
            rejected += 1
            if len(issues) < 500:
                issues.append(ImportIssue(row_number=number, field_name="", message="Строка не записана"))
    return accepted, rejected, issues


def _load(conn, batch_id) -> ImportBatch:
    row = must(one(conn, "SELECT * FROM import_batches WHERE id = %s", (batch_id,)), "Импорт не найден")
    issues = many(
        conn,
        """
        SELECT row_number, field_name, message
        FROM import_issues WHERE batch_id = %s ORDER BY row_number LIMIT 500
        """,
        (batch_id,),
    )
    return ImportBatch(
        id=row["id"],
        domain=row["domain"],
        mode=row["mode"],
        status=row["status"],
        accepted_count=row["accepted_count"],
        rejected_count=row["rejected_count"],
        issues=[ImportIssue(row_number=item["row_number"], field_name=item["field_name"], message=item["message"]) for item in issues],
    )


def _store_issues(conn, batch_id, issues: list[ImportIssue]) -> None:
    conn.execute("DELETE FROM import_issues WHERE batch_id = %s", (batch_id,))
    for issue in issues:
        conn.execute(
            """
            INSERT INTO import_issues (batch_id, row_number, field_name, message)
            VALUES (%s, %s, %s, %s)
            """,
            (batch_id, issue.row_number, issue.field_name, issue.message),
        )


def create_import(conn, user, domain: str, mode: str, dry_run: bool, filename: str, data: bytes) -> ImportBatch:
    user.require("import.run")
    if domain not in DOMAINS:
        raise invalid("import_domain", "Неизвестный раздел импорта")
    if mode not in ("insert_only", "upsert"):
        raise invalid("import_mode", "Режим: insert_only или upsert")
    rows = _parse(filename, data)
    _check_headers(domain, rows)
    file_id = store_bytes(conn, user.id, filename or "import.csv", "application/octet-stream", data)
    batch = conn.execute(
        """
        INSERT INTO import_batches (domain, mode, status, original_name, stored_file_id, created_by)
        VALUES (%s, %s, 'preview', %s, %s, %s)
        RETURNING id
        """,
        (domain, mode, filename or "import.csv", file_id, user.id),
    ).fetchone()
    accepted, rejected, issues = _process(conn, domain, mode, rows, write=not dry_run)
    status = "preview" if dry_run else "applied"
    conn.execute(
        """
        UPDATE import_batches
        SET accepted_count = %s, rejected_count = %s, status = %s,
            applied_at = CASE WHEN %s = 'applied' THEN now() ELSE NULL END
        WHERE id = %s
        """,
        (accepted, rejected, status, status, batch["id"]),
    )
    _store_issues(conn, batch["id"], issues)
    return _load(conn, batch["id"])


def get_import(conn, user, batch_id) -> ImportBatch:
    user.require("import.run")
    return _load(conn, batch_id)


def apply_import(conn, user, batch_id) -> ImportBatch:
    user.require("import.run")
    batch = must(one(conn, "SELECT * FROM import_batches WHERE id = %s", (batch_id,)), "Импорт не найден")
    if batch["status"] == "applied":
        return _load(conn, batch_id)
    data = read_stored(conn, batch["stored_file_id"])
    rows = _parse(batch["original_name"], data)
    _check_headers(batch["domain"], rows)
    accepted, rejected, issues = _process(conn, batch["domain"], batch["mode"], rows, write=True)
    conn.execute(
        """
        UPDATE import_batches
        SET accepted_count = %s, rejected_count = %s, status = 'applied', applied_at = now()
        WHERE id = %s
        """,
        (accepted, rejected, batch_id),
    )
    _store_issues(conn, batch_id, issues)
    return _load(conn, batch_id)


